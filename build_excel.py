#!/usr/bin/env python3
"""
oa_catalog.csv  ->  dissolution_catalog.xlsx

약물별 pH 1.2 / 6.8 용출 문헌(OA)을 엑셀로 정리.
시트 구성:
  - 요약        : BCS 클래스별 통계
  - 약물별집계  : 약물마다 논문 수 / pH 확인 수
  - 전체        : 310편 전체 (PMC 하이퍼링크 포함, 필터/틀고정)
  - BCS_I~IV    : 클래스별 분리 시트

설치:  pip install openpyxl
실행:  python build_excel.py
"""

import csv
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "oa_catalog.csv"
OUT = "dissolution_catalog.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
CLASS_FILLS = {
    "I": "E2EFDA", "II": "FCE4D6", "III": "DDEBF7", "IV": "FFF2CC",
}
YES_FILL = PatternFill("solid", fgColor="C6EFCE")
Q_FILL = PatternFill("solid", fgColor="FFEB9C")
NO_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = ["drug", "bcs_class", "title", "pmcid", "pmc_url", "ph_1_2", "ph_6_8", "notes"]
KOR = ["약물", "BCS", "논문 제목", "PMCID", "PMC 링크", "pH 1.2", "pH 6.8", "비고(시험조건)"]
WIDTHS = [18, 6, 60, 13, 40, 8, 8, 45]


def load():
    with open(SRC, encoding="utf-8-sig") as f:
        r = csv.DictReader(f)
        return [row for row in r]


def style_header(ws, ncol, row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_table(ws, rows, with_link=True):
    # header
    for i, h in enumerate(KOR, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, len(KOR))
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # body
    for ri, row in enumerate(rows, start=2):
        vals = [row[c] for c in COLS]
        for ci, val in enumerate(vals, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(ci in (3, 8)))
            col = COLS[ci - 1]
            if col == "bcs_class":
                cell.fill = PatternFill("solid", fgColor=CLASS_FILLS.get(val, "FFFFFF"))
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if col in ("ph_1_2", "ph_6_8"):
                cell.alignment = Alignment(horizontal="center", vertical="top")
                cell.fill = YES_FILL if val == "Y" else Q_FILL if val == "?" else NO_FILL
            if col == "pmc_url" and with_link and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")
            if col == "pmcid" and val:
                cell.hyperlink = row["pmc_url"]
                cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(KOR))}{len(rows) + 1}"


def main():
    rows = load()
    rows = [r for r in rows if r.get("pmcid", "").startswith("PMC")]
    # 정렬: 클래스 -> 약물 -> 제목
    order = {"I": 0, "II": 1, "III": 2, "IV": 3}
    rows.sort(key=lambda r: (order.get(r["bcs_class"], 9), r["drug"].lower(), r["title"]))

    wb = Workbook()

    # ---- 요약 ----
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "약물별 pH 1.2 / 6.8 buffer 용출 문헌 (오픈액세스) 카탈로그"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"총 {len(rows)}편 · 출처: PubMed Central OA · 수집: 제목/PMCID/pH 메타데이터"
    ws["A2"].font = Font(italic=True, color="808080")

    cls_cnt = Counter(r["bcs_class"] for r in rows)
    both = Counter(r["bcs_class"] for r in rows if r["ph_1_2"] == "Y" and r["ph_6_8"] == "Y")
    drugcnt = defaultdict(set)
    for r in rows:
        drugcnt[r["bcs_class"]].add(r["drug"].lower())

    hdr = ["BCS Class", "설명", "논문 수", "약물 수", "pH1.2&6.8 모두 Y"]
    desc = {"I": "고용해도·고투과", "II": "저용해도·고투과",
            "III": "고용해도·저투과", "IV": "저용해도·저투과"}
    start = 4
    for i, h in enumerate(hdr, start=1):
        ws.cell(row=start, column=i, value=h)
    style_header(ws, len(hdr), row=start)
    rr = start + 1
    for c in ["I", "II", "III", "IV"]:
        ws.cell(row=rr, column=1, value=c).fill = PatternFill("solid", fgColor=CLASS_FILLS[c])
        ws.cell(row=rr, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=rr, column=2, value=desc[c])
        ws.cell(row=rr, column=3, value=cls_cnt.get(c, 0))
        ws.cell(row=rr, column=4, value=len(drugcnt.get(c, [])))
        ws.cell(row=rr, column=5, value=both.get(c, 0))
        rr += 1
    ws.cell(row=rr, column=1, value="합계").font = Font(bold=True)
    ws.cell(row=rr, column=3, value=len(rows)).font = Font(bold=True)
    ws.cell(row=rr, column=4, value=len({r["drug"].lower() for r in rows})).font = Font(bold=True)
    ws.cell(row=rr, column=5, value=sum(both.values())).font = Font(bold=True)
    for i, w in enumerate([12, 18, 10, 10, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---- 약물별 집계 ----
    wsd = wb.create_sheet("약물별집계")
    by = defaultdict(list)
    for r in rows:
        by[(r["bcs_class"], r["drug"])].append(r)
    dh = ["약물", "BCS", "논문 수", "pH1.2 Y", "pH6.8 Y", "둘다 Y"]
    for i, h in enumerate(dh, start=1):
        wsd.cell(row=1, column=i, value=h)
    style_header(wsd, len(dh))
    keys = sorted(by.keys(), key=lambda k: (order.get(k[0], 9), k[1].lower()))
    for ri, k in enumerate(keys, start=2):
        grp = by[k]
        wsd.cell(row=ri, column=1, value=k[1])
        cc = wsd.cell(row=ri, column=2, value=k[0])
        cc.fill = PatternFill("solid", fgColor=CLASS_FILLS.get(k[0], "FFFFFF"))
        cc.alignment = Alignment(horizontal="center")
        wsd.cell(row=ri, column=3, value=len(grp))
        wsd.cell(row=ri, column=4, value=sum(1 for g in grp if g["ph_1_2"] == "Y"))
        wsd.cell(row=ri, column=5, value=sum(1 for g in grp if g["ph_6_8"] == "Y"))
        wsd.cell(row=ri, column=6, value=sum(1 for g in grp if g["ph_1_2"] == "Y" and g["ph_6_8"] == "Y"))
    for i, w in enumerate([20, 6, 9, 9, 9, 9], start=1):
        wsd.column_dimensions[get_column_letter(i)].width = w
    wsd.freeze_panes = "A2"
    wsd.auto_filter.ref = f"A1:F{len(keys) + 1}"

    # ---- 전체 ----
    write_table(wb.create_sheet("전체"), rows)

    # ---- 클래스별 ----
    for c in ["I", "II", "III", "IV"]:
        sub = [r for r in rows if r["bcs_class"] == c]
        write_table(wb.create_sheet(f"BCS_{c}"), sub)

    # ---- 수치데이터 (스니펫 best-effort) ----
    vpath = Path("dissolution_values.csv")
    if vpath.exists():
        wsv = wb.create_sheet("수치데이터")
        vhead = ["약물", "BCS", "pH", "시간", "%release/용해도", "매질", "출처", "비고"]
        vcols = ["drug", "bcs_class", "pH", "time", "percent_released", "medium", "source", "note"]
        for i, h in enumerate(vhead, start=1):
            wsv.cell(row=1, column=i, value=h)
        style_header(wsv, len(vhead))
        with open(vpath, encoding="utf-8-sig") as f:
            vrows = list(csv.DictReader(f))
        vrows.sort(key=lambda r: (order.get(r["bcs_class"], 9), r["drug"].lower()))
        for ri, r in enumerate(vrows, start=2):
            for ci, col in enumerate(vcols, start=1):
                cell = wsv.cell(row=ri, column=ci, value=r[col])
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=(ci == 8))
                if col == "bcs_class":
                    cell.fill = PatternFill("solid", fgColor=CLASS_FILLS.get(r[col], "FFFFFF"))
                    cell.alignment = Alignment(horizontal="center")
                if col == "pH":
                    cell.alignment = Alignment(horizontal="center")
                    if r[col] in ("1.2", "6.8"):
                        cell.fill = YES_FILL
        for i, w in enumerate([18, 6, 7, 10, 16, 26, 26, 40], start=1):
            wsv.column_dimensions[get_column_letter(i)].width = w
        wsv.freeze_panes = "A2"
        wsv.auto_filter.ref = f"A1:H{len(vrows) + 1}"
        # 수치 시트를 요약 다음(2번째)로 이동
        wb.move_sheet("수치데이터", -(len(wb.sheetnames) - 2))

    wb.save(OUT)
    print(f"저장: {OUT}  (총 {len(rows)}편, 시트 {len(wb.sheetnames)}개)")


if __name__ == "__main__":
    main()
